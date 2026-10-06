"""Import only the unfulfilled part of external orders at activation.

Historical stock and settlement are represented by opening balances. Importing an
old shipment, receipt or issue would count it twice, so this module creates draft
documents for remaining work and stores the original progress separately.
"""

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Literal
from uuid import UUID, uuid4
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.core.pagination import paginate
from app.core.security import CurrentUser, require_csrf, require_permission
from app.inventory.quantities import quantity_snapshot

router = APIRouter(prefix="/api/system/carryover", tags=["system"])
Kind = Literal["sales", "purchase", "production", "outsourcing"]
HEADERS = ("原单号", "接续单号", "启用日期", "往来单位编码", "行类型", "物料编码",
           "单位编码", "原数量", "已完成数量", "剩余数量", "单价", "供料方", "历史来源仓库编码")
KIND_NAMES = {"sales": "销售", "purchase": "采购", "production": "生产", "outsourcing": "委外"}
TABLES = {"sales": "sales_orders", "purchase": "purchase_orders",
          "production": "production_orders", "outsourcing": "outsourcing_orders"}
PARTY_TYPES = {"sales": "customer", "purchase": "supplier", "outsourcing": "processor"}
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024


def _template(kind: Kind) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"接续{KIND_NAMES[kind]}"
    sheet.append(HEADERS)
    for column in "ABCDEFGHIJKLM":
        sheet.column_dimensions[column].width = 19
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 25
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _decimal(value: object, row_no: int, label: str, *, allow_zero: bool = False) -> Decimal:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise HTTPException(status_code=422, detail=f"row {row_no}: invalid {label}") from exc
    if not number.is_finite() or number < 0 or (not allow_zero and number == 0) or number > Decimal("99999999999999") or number.as_tuple().exponent < -6:
        raise HTTPException(status_code=422, detail=f"row {row_no}: invalid {label}")
    return number


def _date(value: object, row_no: int) -> date:
    try:
        return value.date() if isinstance(value, datetime) else (
            value if isinstance(value, date) else date.fromisoformat(str(value).strip()))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=f"row {row_no}: invalid effective date") from exc


def _prepare(db: Session, org: UUID, kind: Kind, filename: str, content: bytes) -> list[dict]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="只支持 .xlsx 文件")
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Excel 文件超过 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Excel 文件无效") from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header[:len(HEADERS)]) != HEADERS:
            raise HTTPException(status_code=422, detail="Excel 表头与模板不一致")
        documents: dict[str, dict] = {}
        line_count = 0
        for row_no, row in enumerate(rows, 2):
            values = list(row[:len(HEADERS)]) + [None] * max(0, len(HEADERS) - len(row))
            if not any(value is not None and str(value).strip() for value in values):
                continue
            line_count += 1
            if line_count > MAX_ROWS:
                raise HTTPException(status_code=422, detail=f"maximum {MAX_ROWS} rows exceeded")
            original_no, new_no = (str(values[i]).strip() if values[i] is not None else "" for i in (0, 1))
            if not original_no or not new_no or len(original_no) > 80 or len(new_no) > (55 if kind == "production" else 80):
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid document number")
            effective_date = _date(values[2], row_no)
            if effective_date <= date(1900, 1, 1):
                raise HTTPException(status_code=422, detail=f"row {row_no}: effective date is too early")
            party_code = str(values[3]).strip() if values[3] is not None else ""
            line_kind = {"产出": "output", "物料": "material", "output": "output", "material": "material",
                         "": "output"}.get(str(values[4]).strip() if values[4] is not None else "")
            if line_kind is None or (kind != "outsourcing" and line_kind != "output"):
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid line type")
            item_code = str(values[5]).strip() if values[5] is not None else ""
            unit_code = str(values[6]).strip() if values[6] is not None else ""
            if not item_code or not unit_code:
                raise HTTPException(status_code=422, detail=f"row {row_no}: item and unit are required")
            original = _decimal(values[7], row_no, "original quantity")
            completed = _decimal(values[8], row_no, "completed quantity", allow_zero=True)
            remaining = _decimal(values[9], row_no, "remaining quantity", allow_zero=True)
            if original != completed + remaining:
                raise HTTPException(status_code=422, detail=f"row {row_no}: original != completed + remaining")
            if kind == "production" and any(number != number.to_integral_value() for number in (original, completed, remaining)):
                raise HTTPException(status_code=422, detail=f"row {row_no}: production output must be an integer")
            price = _decimal(values[10], row_no, "unit price", allow_zero=True) if kind in ("sales", "purchase") or (kind == "outsourcing" and line_kind == "output") else None
            supply = {"我方": "self", "加工商": "processor", "self": "self", "processor": "processor"}.get(
                str(values[11]).strip() if values[11] is not None else "")
            if kind == "outsourcing" and line_kind == "material" and supply is None:
                raise HTTPException(status_code=422, detail=f"row {row_no}: material supply party is required")
            warehouse_code = str(values[12]).strip() if values[12] is not None else ""
            warehouse_id = None
            if kind in ("sales", "purchase") and completed > 0:
                warehouse_id = db.execute(text("""
                    SELECT id FROM warehouses
                    WHERE organization_id = :org AND code = :code AND is_active
                """), {"org": org, "code": warehouse_code}).scalar_one_or_none()
                if warehouse_id is None:
                    raise HTTPException(status_code=422, detail=f"row {row_no}: historical source warehouse is required")
            item = db.execute(text("""
                SELECT i.id, i.item_type, i.base_unit_id, u.id AS unit_id
                FROM items i JOIN units u ON u.organization_id = i.organization_id
                WHERE i.organization_id = :org AND i.code = :item AND i.is_active
                  AND u.code = :unit AND u.is_active
            """), {"org": org, "item": item_code, "unit": unit_code}).mappings().first()
            if item is None:
                raise HTTPException(status_code=422, detail=f"row {row_no}: unknown item or unit")
            if kind == "sales" and item["item_type"] != "finished":
                raise HTTPException(status_code=422, detail=f"row {row_no}: sales item must be finished")
            if kind in ("production", "outsourcing") and line_kind == "output" and item["item_type"] not in ("semi_finished", "finished"):
                raise HTTPException(status_code=422, detail=f"row {row_no}: output item type is invalid")
            if kind in ("production", "outsourcing") and item["base_unit_id"] != item["unit_id"]:
                raise HTTPException(status_code=422, detail=f"row {row_no}: use base unit for production/outsourcing")
            factor, original_base = quantity_snapshot(db, org, item["id"], item["unit_id"], original)
            completed_base = quantity_snapshot(db, org, item["id"], item["unit_id"], completed)[1]
            remaining_base = quantity_snapshot(db, org, item["id"], item["unit_id"], remaining)[1]
            if original_base != completed_base + remaining_base:
                raise HTTPException(status_code=422, detail=f"row {row_no}: converted quantities do not balance")
            if price is not None and max(remaining, completed) * price >= Decimal("10000000000000000"):
                raise HTTPException(status_code=422, detail=f"row {row_no}: document amount is too large")
            if kind == "production" and remaining_base > Decimal(9223372036854775807):
                raise HTTPException(status_code=422, detail=f"row {row_no}: production quantity is too large")
            doc = documents.get(original_no)
            if doc is None:
                if kind in PARTY_TYPES:
                    party = db.execute(text("""
                        SELECT p.id, p.name FROM parties p JOIN party_types t ON t.party_id = p.id
                        WHERE p.organization_id = :org AND p.code = :code AND p.is_active
                          AND t.party_type = :account
                    """), {"org": org, "code": party_code, "account": PARTY_TYPES[kind]}).mappings().first()
                    if party is None:
                        raise HTTPException(status_code=422, detail=f"row {row_no}: unknown party/account")
                    party_id, party_name = party["id"], party["name"]
                else:
                    if party_code:
                        raise HTTPException(status_code=422, detail=f"row {row_no}: production has no party")
                    party_id, party_name = None, None
                exists = db.execute(text("""
                    SELECT 1 FROM carryover_documents
                    WHERE organization_id = :org AND business_kind = :kind AND original_document_no = :original
                """), {"org": org, "kind": kind, "original": original_no}).scalar_one_or_none()
                if exists:
                    raise HTTPException(status_code=409, detail=f"row {row_no}: original document already imported")
                doc = {"original_document_no": original_no, "document_no": new_no,
                       "effective_date": effective_date, "party_id": party_id,
                       "party_code": party_code, "party_name": party_name, "lines": []}
                documents[original_no] = doc
            elif (doc["document_no"], doc["effective_date"], doc["party_code"]) != (new_no, effective_date, party_code):
                raise HTTPException(status_code=422, detail=f"row {row_no}: document header differs from earlier row")
            doc["lines"].append({"row_no": row_no, "line_kind": line_kind, "item_id": item["id"],
                "item_code": item_code, "unit_id": item["unit_id"], "unit_code": unit_code,
                "factor": factor, "original_base": original_base,
                "completed_base": completed_base, "remaining_base": remaining_base,
                "completed": completed, "remaining": remaining, "price": price,
                "supply_party": supply, "warehouse_id": warehouse_id,
                "warehouse_code": warehouse_code})
        if not documents:
            raise HTTPException(status_code=422, detail="Excel 中没有单据数据")
        new_numbers = [doc["document_no"] for doc in documents.values()]
        if len(new_numbers) != len(set(new_numbers)):
            raise HTTPException(status_code=422, detail="接续单据号重复")
        for doc in documents.values():
            positive_outputs = [line for line in doc["lines"] if line["line_kind"] == "output" and line["remaining"] > 0]
            if not positive_outputs:
                raise HTTPException(status_code=422, detail=f"{doc['original_document_no']}: no remaining output")
            if kind in ("production", "outsourcing"):
                keys = [(line["item_id"], line["line_kind"], line["supply_party"]) for line in doc["lines"]]
                if len(keys) != len(set(keys)):
                    raise HTTPException(status_code=422, detail=f"{doc['original_document_no']}: duplicate item/line type")
            target_exists = db.execute(text(f"""
                SELECT 1 FROM {TABLES[kind]}
                WHERE organization_id = :org AND document_no = :number
            """), {"org": org, "number": doc["document_no"]}).scalar_one_or_none()
            if target_exists:
                raise HTTPException(status_code=409, detail=f"{doc['document_no']}: document number already exists")
        return list(documents.values())
    finally:
        workbook.close()


def _insert_document(db: Session, org: UUID, actor: UUID, kind: Kind, doc: dict) -> tuple[UUID, UUID | None]:
    common = {"org": org, "number": doc["document_no"], "date": doc["effective_date"],
              "actor": actor, "remark": f"期初接续；原单号：{doc['original_document_no']}。仅承接剩余待履约数量。"}
    active = [line for line in doc["lines"] if line["remaining"] > 0]
    if kind == "sales":
        doc_id = db.execute(text("""
            INSERT INTO sales_orders (organization_id, document_no, customer_id, document_date, remark, created_by)
            VALUES (:org, :number, :party, :date, :remark, :actor) RETURNING id
        """), {**common, "party": doc["party_id"]}).scalar_one()
        for line in active:
            amount = (line["remaining"] * line["price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            db.execute(text("""
                INSERT INTO sales_order_lines (sales_order_id, item_id, quantity, unit_id,
                    conversion_factor, quantity_base, unit_price, amount)
                VALUES (:doc, :item, :quantity, :unit, :factor, :base, :price, :amount)
            """), {"doc": doc_id, "item": line["item_id"], "quantity": line["remaining"],
                    "unit": line["unit_id"], "factor": line["factor"], "base": line["remaining_base"],
                    "price": line["price"], "amount": amount})
        return doc_id, None
    if kind == "purchase":
        doc_id = db.execute(text("""
            INSERT INTO purchase_orders (organization_id, document_no, supplier_id, document_date, remark, created_by)
            VALUES (:org, :number, :party, :date, :remark, :actor) RETURNING id
        """), {**common, "party": doc["party_id"]}).scalar_one()
        for line in active:
            amount = (line["remaining"] * line["price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            db.execute(text("""
                INSERT INTO purchase_order_lines (purchase_order_id, item_id, quantity, unit_id,
                    conversion_factor, quantity_base, unit_price, amount)
                VALUES (:doc, :item, :quantity, :unit, :factor, :base, :price, :amount)
            """), {"doc": doc_id, "item": line["item_id"], "quantity": line["remaining"],
                    "unit": line["unit_id"], "factor": line["factor"], "base": line["remaining_base"],
                    "price": line["price"], "amount": amount})
        return doc_id, None
    if kind == "production":
        plan_no = doc["document_no"] + "-PLAN"
        plan_id = db.execute(text("""
            INSERT INTO production_plans (organization_id, document_no, document_date, remark, created_by)
            VALUES (:org, :number, :date, :remark, :actor) RETURNING id
        """), {**common, "number": plan_no}).scalar_one()
        plan_lines = {}
        for line in active:
            plan_lines[line["row_no"]] = db.execute(text("""
                INSERT INTO production_plan_lines (production_plan_id, item_id, planned_quantity_base)
                VALUES (:plan, :item, :quantity) RETURNING id
            """), {"plan": plan_id, "item": line["item_id"],
                    "quantity": int(line["remaining_base"])}).scalar_one()
        doc_id = db.execute(text("""
            INSERT INTO production_orders (organization_id, document_no, production_plan_id,
                document_date, remark, created_by)
            VALUES (:org, :number, :plan, :date, :remark, :actor) RETURNING id
        """), {**common, "plan": plan_id}).scalar_one()
        for line in active:
            db.execute(text("""
                INSERT INTO production_order_outputs (production_order_id, production_plan_line_id,
                    item_id, planned_quantity_base)
                VALUES (:doc, :plan_line, :item, :quantity)
            """), {"doc": doc_id, "plan_line": plan_lines[line["row_no"]],
                    "item": line["item_id"], "quantity": int(line["remaining_base"])})
        return doc_id, plan_id
    doc_id = db.execute(text("""
        INSERT INTO outsourcing_orders (organization_id, document_no, processor_id,
            document_date, remark, created_by)
        VALUES (:org, :number, :party, :date, :remark, :actor) RETURNING id
    """), {**common, "party": doc["party_id"]}).scalar_one()
    for line in active:
        if line["line_kind"] == "material":
            db.execute(text("""
                INSERT INTO outsourcing_material_lines (outsourcing_order_id, item_id,
                    supply_party, expected_quantity_base)
                VALUES (:doc, :item, :supply, :quantity)
            """), {"doc": doc_id, "item": line["item_id"], "supply": line["supply_party"],
                    "quantity": line["remaining_base"]})
        else:
            db.execute(text("""
                INSERT INTO outsourcing_output_lines (outsourcing_order_id, item_id,
                    expected_quantity_base, unit_price)
                VALUES (:doc, :item, :quantity, :price)
            """), {"doc": doc_id, "item": line["item_id"],
                    "quantity": line["remaining_base"], "price": line["price"]})
    return doc_id, None


def _insert_historical_sources(
    db: Session, org: UUID, actor: UUID, kind: Kind, doc: dict,
) -> tuple[UUID | None, dict[int, tuple[UUID, UUID]]]:
    """Create source references for later returns; no stock or amount is posted."""
    completed = [line for line in doc["lines"] if line["completed"] > 0]
    if kind not in ("sales", "purchase") or not completed:
        return None, {}
    source_date = doc["effective_date"] - timedelta(days=1)
    prefix = "SO" if kind == "sales" else "PO"
    historical_no = f"HIST-{prefix}-{uuid4().hex[:16].upper()}"
    remark = f"期初历史来源；原单号：{doc['original_document_no']}。仅供退货与补发关联，不得重放库存及往来。"
    if kind == "sales":
        order_id = db.execute(text("""
            INSERT INTO sales_orders (organization_id, document_no, customer_id,
                document_date, status, remark, created_by)
            VALUES (:org, :number, :party, :date, 'approved', :remark, :actor) RETURNING id
        """), {"org": org, "number": historical_no, "party": doc["party_id"],
                "date": source_date, "remark": remark, "actor": actor}).scalar_one()
    else:
        order_id = db.execute(text("""
            INSERT INTO purchase_orders (organization_id, document_no, supplier_id,
                document_date, status, remark, created_by)
            VALUES (:org, :number, :party, :date, 'open', :remark, :actor) RETURNING id
        """), {"org": org, "number": historical_no, "party": doc["party_id"],
                "date": source_date, "remark": remark, "actor": actor}).scalar_one()
    references: dict[int, tuple[UUID, UUID]] = {}
    for line in completed:
        if kind == "sales":
            amount = (line["completed"] * line["price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            order_line_id = db.execute(text("""
                INSERT INTO sales_order_lines (sales_order_id, item_id, quantity, unit_id,
                    conversion_factor, quantity_base, unit_price, amount)
                VALUES (:order, :item, :quantity, :unit, :factor, :base, :price, :amount)
                RETURNING id
            """), {"order": order_id, "item": line["item_id"], "quantity": line["completed"],
                    "unit": line["unit_id"], "factor": line["factor"], "base": line["completed_base"],
                    "price": line["price"], "amount": amount}).scalar_one()
            source_id = db.execute(text("""
                INSERT INTO sales_shipments (organization_id, document_no, sales_order_id,
                    warehouse_id, document_date, shipment_type, status, remark,
                    created_by, posted_at, is_opening_reference)
                VALUES (:org, :number, :order, :warehouse, :date, 'normal', 'posted',
                    :remark, :actor, CAST(:effective AS timestamptz) - interval '1 second', true)
                RETURNING id
            """), {"org": org, "number": f"HIST-SH-{doc['original_document_no'][:30]}-{uuid4().hex[:6].upper()}",
                    "order": order_id, "warehouse": line["warehouse_id"],
                    "date": source_date, "effective": doc["effective_date"],
                    "remark": remark, "actor": actor}).scalar_one()
            source_line_id = db.execute(text("""
                INSERT INTO sales_shipment_lines (shipment_id, sales_order_line_id,
                    item_id, quantity, unit_id, conversion_factor, quantity_base)
                VALUES (:source, :order_line, :item, :quantity, :unit, :factor, :base)
                RETURNING id
            """), {"source": source_id, "order_line": order_line_id,
                    "item": line["item_id"], "quantity": line["completed"],
                    "unit": line["unit_id"], "factor": line["factor"],
                    "base": line["completed_base"]}).scalar_one()
        else:
            amount = (line["completed"] * line["price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            order_line_id = db.execute(text("""
                INSERT INTO purchase_order_lines (purchase_order_id, item_id, quantity,
                    unit_id, conversion_factor, quantity_base, unit_price, amount)
                VALUES (:order, :item, :quantity, :unit, :factor, :base, :price, :amount)
                RETURNING id
            """), {"order": order_id, "item": line["item_id"], "quantity": line["completed"],
                    "unit": line["unit_id"], "factor": line["factor"], "base": line["completed_base"],
                    "price": line["price"], "amount": amount}).scalar_one()
            source_id = db.execute(text("""
                INSERT INTO purchase_receipts (organization_id, document_no, purchase_order_id,
                    warehouse_id, document_date, status, remark, created_by, posted_at,
                    is_opening_reference)
                VALUES (:org, :number, :order, :warehouse, :date, 'posted', :remark,
                    :actor, CAST(:effective AS timestamptz) - interval '1 second', true)
                RETURNING id
            """), {"org": org, "number": f"HIST-PR-{doc['original_document_no'][:30]}-{uuid4().hex[:6].upper()}",
                    "order": order_id, "warehouse": line["warehouse_id"],
                    "date": source_date, "effective": doc["effective_date"],
                    "remark": remark, "actor": actor}).scalar_one()
            source_line_id = db.execute(text("""
                INSERT INTO purchase_receipt_lines (receipt_id, purchase_order_line_id,
                    item_id, quantity, unit_id, conversion_factor, quantity_base,
                    unit_price, amount)
                VALUES (:source, :order_line, :item, :quantity, :unit, :factor,
                    :base, :price, :amount)
                RETURNING id
            """), {"source": source_id, "order_line": order_line_id,
                    "item": line["item_id"], "quantity": line["completed"],
                    "unit": line["unit_id"], "factor": line["factor"],
                    "base": line["completed_base"], "price": line["price"],
                    "amount": amount}).scalar_one()
        references[line["row_no"]] = (source_id, source_line_id)
    return order_id, references


def _import(db: Session, org: UUID, actor: UUID, kind: Kind, documents: list[dict]) -> dict:
    try:
        for doc in documents:
            doc_id, plan_id = _insert_document(db, org, actor, kind, doc)
            historical_order_id, historical_sources = _insert_historical_sources(
                db, org, actor, kind, doc)
            carry_id = db.execute(text("""
                INSERT INTO carryover_documents (organization_id, business_kind,
                    original_document_no, document_no, effective_date, party_id,
                    document_id, related_plan_id, historical_order_id, imported_by)
                VALUES (:org, :kind, :original, :number, :date, :party,
                    :document_id, :plan_id, :historical_order_id, :actor) RETURNING id
            """), {"org": org, "kind": kind, "original": doc["original_document_no"],
                    "number": doc["document_no"], "date": doc["effective_date"],
                    "party": doc["party_id"], "document_id": doc_id,
                    "plan_id": plan_id, "historical_order_id": historical_order_id,
                    "actor": actor}).scalar_one()
            for line in doc["lines"]:
                source_id, source_line_id = historical_sources.get(line["row_no"], (None, None))
                db.execute(text("""
                    INSERT INTO carryover_lines (carryover_document_id, source_row_no,
                        line_kind, item_id, unit_id, original_quantity_base,
                        completed_quantity_base, remaining_quantity_base, unit_price,
                        supply_party, historical_source_id, historical_source_line_id)
                    VALUES (:carry, :row, :line_kind, :item, :unit, :original,
                        :completed, :remaining, :price, :supply, :source, :source_line)
                """), {"carry": carry_id, "row": line["row_no"], "line_kind": line["line_kind"],
                        "item": line["item_id"], "unit": line["unit_id"],
                        "original": line["original_base"], "completed": line["completed_base"],
                        "remaining": line["remaining_base"], "price": line["price"],
                        "supply": line["supply_party"], "source": source_id,
                        "source_line": source_line_id})
            db.execute(text("""
                INSERT INTO audit_logs (organization_id, actor_id, action_code, document_type, document_id)
                VALUES (:org, :actor, 'carryover.import', 'carryover', :id)
            """), {"org": org, "actor": actor, "id": carry_id})
        db.commit()
        return {"imported_documents": len(documents),
                "imported_lines": sum(len(doc["lines"]) for doc in documents)}
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="接续数据与现有数据冲突，没有导入任何单据") from exc


@router.get("/template/{kind}")
def template(kind: Kind,
             user: CurrentUser = Depends(require_permission("system.initialization", "view"))) -> StreamingResponse:
    return StreamingResponse(BytesIO(_template(kind)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="carryover-{kind}-template.xlsx"'})


@router.post("/preview/{kind}", dependencies=[Depends(require_csrf)])
def preview(kind: Kind, file: UploadFile = File(...),
            user: CurrentUser = Depends(require_permission("system.initialization", "view")),
            db: Session = Depends(get_db)) -> list[dict]:
    documents = _prepare(db, user.organization_id, kind, file.filename or "",
                         file.file.read(MAX_BYTES + 1))
    return [{"original_document_no": doc["original_document_no"],
             "document_no": doc["document_no"], "effective_date": doc["effective_date"],
             "party_code": doc["party_code"], "line_count": len(doc["lines"]),
             "remaining_lines": sum(line["remaining"] > 0 for line in doc["lines"])}
            for doc in documents]


@router.post("/import/{kind}", dependencies=[Depends(require_csrf)])
def import_documents(kind: Kind, file: UploadFile = File(...),
                     user: CurrentUser = Depends(require_permission("system.initialization", "create")),
                     db: Session = Depends(get_db)) -> dict:
    documents = _prepare(db, user.organization_id, kind, file.filename or "",
                         file.file.read(MAX_BYTES + 1))
    return _import(db, user.organization_id, user.id, kind, documents)


@router.get("/documents")
def list_documents(response: Response,
                   kind: Kind | None = None,
                   limit: int = Query(default=100, ge=1, le=500), offset: int = Query(default=0, ge=0),
                   user: CurrentUser = Depends(require_permission("system.initialization", "view")),
                   db: Session = Depends(get_db)) -> list[dict]:
    query = """
        SELECT c.id, c.business_kind, c.original_document_no, c.document_no,
               c.effective_date, c.document_id, c.related_plan_id,
               c.historical_order_id, c.imported_at,
               count(l.id) AS line_count,
               COALESCE(sum(l.remaining_quantity_base), 0) AS remaining_quantity_base
        FROM carryover_documents c LEFT JOIN carryover_lines l ON l.carryover_document_id = c.id
        WHERE c.organization_id = :org
          AND (CAST(:kind AS text) IS NULL OR c.business_kind = :kind)
        GROUP BY c.id ORDER BY c.imported_at DESC, c.id DESC
    """
    return paginate(db, query, {"org": user.organization_id, "kind": kind},
                    limit=limit, offset=offset, response=response)


@router.get("/documents/{carry_id}")
def get_document(carry_id: UUID,
                 user: CurrentUser = Depends(require_permission("system.initialization", "view")),
                 db: Session = Depends(get_db)) -> dict:
    document = db.execute(text("""
        SELECT id, business_kind, original_document_no, document_no,
               effective_date, document_id, related_plan_id, historical_order_id,
               imported_at
        FROM carryover_documents
        WHERE id = :id AND organization_id = :org
    """), {"id": carry_id, "org": user.organization_id}).mappings().first()
    if document is None:
        raise HTTPException(status_code=404, detail="接续单据不存在")
    result = dict(document)
    result["lines"] = [dict(row) for row in db.execute(text("""
        SELECT l.source_row_no, l.line_kind, i.code AS item_code, i.name AS item_name,
               u.code AS unit_code, l.original_quantity_base,
               l.completed_quantity_base, l.remaining_quantity_base,
               l.unit_price, l.supply_party, l.historical_source_id,
               l.historical_source_line_id,
               COALESCE(s.document_no, p.document_no) AS historical_source_no
        FROM carryover_lines l
        JOIN items i ON i.id = l.item_id JOIN units u ON u.id = l.unit_id
        LEFT JOIN sales_shipments s ON s.id = l.historical_source_id
        LEFT JOIN purchase_receipts p ON p.id = l.historical_source_id
        WHERE l.carryover_document_id = :id
        ORDER BY l.source_row_no
    """), {"id": carry_id}).mappings()]
    return result
