"""Excel import and posting of opening stock quantities."""

from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import UUID, uuid4
from zipfile import BadZipFile

from fastapi import HTTPException, Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.pagination import paginate
from app.inventory.posting import StockChange, post_stock_changes
from app.inventory.quantities import quantity_snapshot

HEADERS = ("仓库编码", "物料编码", "单位编码", "数量")
MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024


def template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "期初库存"
    sheet.append(HEADERS)
    for column, width in {"A": 18, "B": 22, "C": 16, "D": 18}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _parse_rows(content: bytes) -> list[tuple[int, str, str, str, Decimal]]:
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Excel 文件超过 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Excel 文件无效") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header[:4]) != HEADERS:
            raise HTTPException(status_code=422, detail="Excel 表头与模板不一致")
        result: list[tuple[int, str, str, str, Decimal]] = []
        for row_no, row in enumerate(rows, start=2):
            values = row[:4]
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(result) >= MAX_ROWS:
                raise HTTPException(status_code=422, detail=f"maximum {MAX_ROWS} rows exceeded")
            if len(values) < 4 or any(value is None or not str(value).strip() for value in values):
                raise HTTPException(status_code=422, detail=f"row {row_no}: required cell is empty")
            warehouse_code, item_code, unit_code = (str(value).strip() for value in values[:3])
            try:
                quantity = Decimal(str(values[3]))
            except (InvalidOperation, ValueError) as exc:
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid quantity") from exc
            if not quantity.is_finite() or quantity <= 0 or quantity > Decimal("99999999999999"):
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid quantity")
            result.append((row_no, warehouse_code, item_code, unit_code, quantity))
        if not result:
            raise HTTPException(status_code=422, detail="Excel 中没有库存数据行")
        return result
    finally:
        workbook.close()


def _get_document(db: Session, organization_id: UUID, document_id: UUID, *, lock: bool = False) -> dict | None:
    suffix = " FOR UPDATE" if lock else ""
    row = db.execute(
        text("""
            SELECT id, organization_id, document_no, effective_date, import_batch_no,
                   status, remark, created_at, updated_at
            FROM opening_stock_docs WHERE id = :id AND organization_id = :org
        """ + suffix),
        {"id": document_id, "org": organization_id},
    ).mappings().first()
    if row is None:
        return None
    result = dict(row)
    result["lines"] = [dict(line) for line in db.execute(
        text("""
            SELECT id, warehouse_id, item_id, unit_id, quantity, conversion_factor, quantity_base
            FROM opening_stock_lines WHERE opening_doc_id = :id ORDER BY sort_order, id
        """),
        {"id": document_id},
    ).mappings()]
    return result


def get_document(db: Session, organization_id: UUID, document_id: UUID) -> dict:
    result = _get_document(db, organization_id, document_id)
    if result is None:
        raise HTTPException(status_code=404, detail="期初库存单不存在")
    return result


def list_documents(db: Session, organization_id: UUID, *, status: str | None,
                   limit: int, offset: int, response: Response) -> list[dict]:
    query = """
        SELECT id FROM opening_stock_docs
        WHERE organization_id = :org
          AND (CAST(:status AS text) IS NULL OR status = :status)
        ORDER BY created_at DESC, document_no DESC
    """
    ids = paginate(db, query, {"org": organization_id, "status": status},
                   limit=limit, offset=offset, response=response)
    return [get_document(db, organization_id, row["id"]) for row in ids]


def prepare_lines(db: Session, organization_id: UUID, filename: str, content: bytes) -> list[dict]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="只支持 .xlsx 文件")
    parsed = _parse_rows(content)
    prepared: list[dict] = []
    seen: set[tuple[UUID, UUID]] = set()
    for row_no, warehouse_code, item_code, unit_code, quantity in parsed:
        warehouse_id = db.execute(
            text("""
                SELECT id FROM warehouses
                WHERE organization_id = :org AND code = :code AND is_active
            """),
            {"org": organization_id, "code": warehouse_code},
        ).scalar_one_or_none()
        item_id = db.execute(
            text("""
                SELECT id FROM items
                WHERE organization_id = :org AND code = :code AND is_active
            """),
            {"org": organization_id, "code": item_code},
        ).scalar_one_or_none()
        unit_id = db.execute(
            text("""
                SELECT id FROM units
                WHERE organization_id = :org AND code = :code AND is_active
            """),
            {"org": organization_id, "code": unit_code},
        ).scalar_one_or_none()
        if None in (warehouse_id, item_id, unit_id):
            raise HTTPException(status_code=422, detail=f"row {row_no}: unknown or inactive master data")
        key = (warehouse_id, item_id)
        if key in seen:
            raise HTTPException(status_code=422, detail=f"row {row_no}: duplicate warehouse and item")
        seen.add(key)
        try:
            factor, quantity_base = quantity_snapshot(db, organization_id, item_id, unit_id, quantity)
        except HTTPException as exc:
            raise HTTPException(status_code=422, detail=f"row {row_no}: {exc.detail}") from exc
        prepared.append({
            "row_no": row_no, "warehouse_code": warehouse_code,
            "item_code": item_code, "unit_code": unit_code,
            "warehouse_id": warehouse_id, "item_id": item_id, "unit_id": unit_id,
            "quantity": quantity, "factor": factor, "quantity_base": quantity_base,
        })
    return prepared


def import_document(
    db: Session, *, organization_id: UUID, actor_id: UUID,
    effective_date: date, filename: str, content: bytes,
) -> dict:
    prepared = prepare_lines(db, organization_id, filename, content)

    document_no = "OPEN-" + uuid4().hex[:16].upper()
    document_id = db.execute(
        text("""
            INSERT INTO opening_stock_docs (
                organization_id, document_no, effective_date, import_batch_no, created_by
            ) VALUES (:org, :document_no, :effective_date, :filename, :actor_id)
            RETURNING id
        """),
        {
            "org": organization_id, "document_no": document_no,
            "effective_date": effective_date, "filename": filename[:200], "actor_id": actor_id,
        },
    ).scalar_one()
    for index, line in enumerate(prepared):
        db.execute(
            text("""
                INSERT INTO opening_stock_lines (
                    opening_doc_id, warehouse_id, item_id, quantity,
                    unit_id, conversion_factor, quantity_base, sort_order
                ) VALUES (
                    :document_id, :warehouse_id, :item_id, :quantity,
                    :unit_id, :factor, :quantity_base, :sort_order
                )
            """),
            {"document_id": document_id, **line, "sort_order": index},
        )
    db.execute(
        text("""
            INSERT INTO audit_logs (organization_id, actor_id, action_code, document_type, document_id)
            VALUES (:org, :actor_id, 'opening_stock.import', 'opening_stock', :id)
        """),
        {"org": organization_id, "actor_id": actor_id, "id": document_id},
    )
    result = get_document(db, organization_id, document_id)
    db.commit()
    return result


def post_document(db: Session, organization_id: UUID, actor_id: UUID, document_id: UUID) -> dict:
    try:
        document = _get_document(db, organization_id, document_id, lock=True)
        if document is None:
            raise HTTPException(status_code=404, detail="期初库存单不存在")
        if document["status"] == "posted":
            return document
        if document["status"] != "draft":
            raise HTTPException(status_code=409, detail="期初库存单不能过账")
        changes = [StockChange(
            warehouse_id=line["warehouse_id"],
            item_id=line["item_id"],
            quantity_delta_base=line["quantity_base"],
            source_type="opening_stock",
            source_id=document_id,
            source_line_id=line["id"],
            movement_kind="in",
        ) for line in document["lines"]]
        post_stock_changes(db, organization_id=organization_id, actor_id=actor_id, changes=changes)
        for line in document["lines"]:
            current_quantity = db.execute(
                text("""
                    SELECT quantity_base FROM stock_balances
                    WHERE organization_id = :org AND warehouse_id = :warehouse_id
                      AND item_id = :item_id
                """),
                {
                    "org": organization_id, "warehouse_id": line["warehouse_id"],
                    "item_id": line["item_id"],
                },
            ).scalar_one()
            if current_quantity != line["quantity_base"]:
                raise HTTPException(status_code=409, detail="该仓库与物料已有库存流水，不能再作为期初导入")
            prior = db.execute(
                text("""
                    SELECT 1 FROM stock_movements
                    WHERE organization_id = :org AND warehouse_id = :warehouse_id
                      AND item_id = :item_id
                      AND (source_type <> 'opening_stock' OR source_id <> :document_id)
                    LIMIT 1
                """),
                {
                    "org": organization_id, "warehouse_id": line["warehouse_id"],
                    "item_id": line["item_id"], "document_id": document_id,
                },
            ).scalar_one_or_none()
            if prior:
                raise HTTPException(status_code=409, detail="该仓库与物料的期初库存已存在")
        db.execute(
            text("UPDATE opening_stock_docs SET status = 'posted', posted_at = now() WHERE id = :id"),
            {"id": document_id},
        )
        db.execute(
            text("""
                INSERT INTO audit_logs (organization_id, actor_id, action_code, document_type, document_id)
                VALUES (:org, :actor_id, 'opening_stock.post', 'opening_stock', :id)
            """),
            {"org": organization_id, "actor_id": actor_id, "id": document_id},
        )
        db.commit()
        document["status"] = "posted"
        return document
    except Exception:
        db.rollback()
        raise


def reverse_document(db: Session, organization_id: UUID, actor_id: UUID,
                     document_id: UUID, reason: str) -> dict:
    document = _get_document(db, organization_id, document_id, lock=True)
    if document is None:
        raise HTTPException(status_code=404, detail="期初库存单不存在")
    if document["status"] == "reversed":
        return document
    if document["status"] != "posted":
        raise HTTPException(status_code=409, detail="只有已过账的期初库存可以冲销")
    changes = []
    for line in document["lines"]:
        original_id = db.execute(text("""
            SELECT id FROM stock_movements
            WHERE organization_id = :org AND source_type = 'opening_stock'
              AND source_id = :doc_id AND source_line_id = :line_id
        """), {"org": organization_id, "doc_id": document_id,
               "line_id": line["id"]}).scalar_one()
        changes.append(StockChange(
            warehouse_id=line["warehouse_id"], item_id=line["item_id"],
            quantity_delta_base=-line["quantity_base"],
            source_type="opening_stock_reversal", source_id=document_id,
            source_line_id=line["id"], movement_kind="opening_reverse_out",
            reversal_of_id=original_id))
    post_stock_changes(db, organization_id=organization_id,
                       actor_id=actor_id, changes=changes)
    db.execute(text("""
        UPDATE opening_stock_docs SET status = 'reversed'
        WHERE id = :id AND organization_id = :org
    """), {"id": document_id, "org": organization_id})
    db.execute(text("""
        INSERT INTO audit_logs (organization_id, actor_id, action_code,
                                document_type, document_id, reason)
        VALUES (:org, :actor, 'opening_stock.reverse',
                'opening_stock', :id, :reason)
    """), {"org": organization_id, "actor": actor_id,
           "id": document_id, "reason": reason.strip()})
    result = _get_document(db, organization_id, document_id)
    db.commit()
    assert result is not None
    return result
