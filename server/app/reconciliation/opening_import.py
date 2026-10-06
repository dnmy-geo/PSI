"""Validated, atomic Excel import of party opening balances."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

HEADERS = ("往来单位编码", "账户类型", "启用日期", "期初金额")
ACCOUNT_TYPES = {"客户": "customer", "供应商": "supplier", "加工商": "processor",
                 "customer": "customer", "supplier": "supplier", "processor": "processor"}
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024


def template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "期初往来余额"
    sheet.append(HEADERS)
    for column, width in {"A": 24, "B": 18, "C": 18, "D": 18}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def prepare(db: Session, organization_id: UUID, filename: str, content: bytes) -> list[dict]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="只支持 .xlsx 文件")
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Excel 文件超过 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Excel 文件无效") from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header[:4]) != HEADERS:
            raise HTTPException(status_code=422, detail="Excel 表头与模板不一致")
        prepared = []
        seen = set()
        for row_no, row in enumerate(rows, 2):
            cells = row[:4]
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            if len(prepared) >= MAX_ROWS:
                raise HTTPException(status_code=422, detail=f"maximum {MAX_ROWS} rows exceeded")
            if len(cells) < 4 or any(value is None or not str(value).strip() for value in cells):
                raise HTTPException(status_code=422, detail=f"row {row_no}: required cell is empty")
            code = str(cells[0]).strip()
            account = ACCOUNT_TYPES.get(str(cells[1]).strip())
            if account is None:
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid account type")
            value = cells[2]
            try:
                effective_date = value.date() if isinstance(value, datetime) else (
                    value if isinstance(value, date) else date.fromisoformat(str(value).strip()))
            except ValueError as exc:
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid date") from exc
            try:
                amount = Decimal(str(cells[3]))
            except (InvalidOperation, ValueError) as exc:
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid amount") from exc
            if not amount.is_finite() or amount < 0 or amount > Decimal("9999999999999999.99") or amount.as_tuple().exponent < -2:
                raise HTTPException(status_code=422, detail=f"row {row_no}: invalid amount")
            party = db.execute(text("""
                SELECT p.id, p.name FROM parties p JOIN party_types t ON t.party_id = p.id
                WHERE p.organization_id = :org AND p.code = :code AND p.is_active
                  AND t.party_type = :account
            """), {"org": organization_id, "code": code, "account": account}).mappings().first()
            if party is None:
                raise HTTPException(status_code=422, detail=f"row {row_no}: unknown or inactive party/account")
            key = (party["id"], account, effective_date)
            if key in seen:
                raise HTTPException(status_code=422, detail=f"row {row_no}: duplicate party/account/date")
            seen.add(key)
            exists = db.execute(text("""
                SELECT 1 FROM party_opening_balances
                WHERE organization_id = :org AND party_id = :party
                  AND account_type = :account AND effective_date = :date
            """), {"org": organization_id, "party": party["id"], "account": account,
                   "date": effective_date}).scalar_one_or_none()
            if exists:
                raise HTTPException(status_code=409, detail=f"row {row_no}: opening balance already exists")
            prepared.append({"row_no": row_no, "party_id": party["id"], "party_code": code,
                             "party_name": party["name"], "account_type": account,
                             "effective_date": effective_date, "amount": amount})
        if not prepared:
            raise HTTPException(status_code=422, detail="Excel 中没有期初往来数据")
        return prepared
    finally:
        workbook.close()


def import_rows(db: Session, organization_id: UUID, actor_id: UUID, rows: list[dict]) -> int:
    try:
        for row in rows:
            doc_id = db.execute(text("""
                INSERT INTO party_opening_balances
                    (organization_id, party_id, account_type, direction, effective_date, amount)
                VALUES (:org, :party, :account, :direction, :date, :amount)
                RETURNING id
            """), {"org": organization_id, "party": row["party_id"],
                   "account": row["account_type"],
                   "direction": "receivable" if row["account_type"] == "customer" else "payable",
                   "date": row["effective_date"], "amount": row["amount"]}).scalar_one()
            db.execute(text("""
                INSERT INTO audit_logs (organization_id, actor_id, action_code, document_type, document_id)
                VALUES (:org, :actor, 'opening_balance.import', 'party_opening_balance', :id)
            """), {"org": organization_id, "actor": actor_id, "id": doc_id})
        db.commit()
        return len(rows)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="期初余额已存在，没有导入任何数据") from exc
